"""
Unit tests for core.export.export_workbook.
"""

from __future__ import annotations

import unittest
from io import BytesIO

import openpyxl
import pandas as pd

from core.export import export_workbook


class TestExportWorkbook(unittest.TestCase):
    def test_sheet_names_and_content_round_trip(self):
        results = pd.DataFrame(
            {"Date": ["2026-08-03"], "SAP": [100.0], "Cadency": [100.0], "Status": ["RECONCILED"]}
        )
        cadency_detail = pd.DataFrame({"Amount": [100.0]})

        raw = export_workbook([("Results", results), ("Cadency - Raw", cadency_detail)])
        wb = openpyxl.load_workbook(BytesIO(raw))

        self.assertEqual(wb.sheetnames, ["Results", "Cadency - Raw"])
        ws = wb["Results"]
        self.assertEqual([c.value for c in ws[1]], ["Date", "SAP", "Cadency", "Status"])
        self.assertEqual([c.value for c in ws[2]], ["2026-08-03", 100.0, 100.0, "RECONCILED"])

    def test_status_rows_are_color_filled(self):
        results = pd.DataFrame({"Status": ["RECONCILED", "DOES NOT RECONCILE", "MANUAL REVIEW"]})
        raw = export_workbook([("Results", results)])
        wb = openpyxl.load_workbook(BytesIO(raw))
        ws = wb["Results"]

        self.assertEqual(ws.cell(row=2, column=1).fill.fgColor.rgb, "00D4EDDA")
        self.assertEqual(ws.cell(row=3, column=1).fill.fgColor.rgb, "00F8D7DA")
        self.assertEqual(ws.cell(row=4, column=1).fill.fgColor.rgb, "00FFF3CD")

    def test_sheet_name_truncated_to_excel_limit(self):
        long_name = "A" * 40
        raw = export_workbook([(long_name, pd.DataFrame({"x": [1]}))])
        wb = openpyxl.load_workbook(BytesIO(raw))
        self.assertEqual(wb.sheetnames, [long_name[:31]])


if __name__ == "__main__":
    unittest.main()
