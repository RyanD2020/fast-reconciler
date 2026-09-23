"""
Excel export -- packages a reconciliation run into a single .xlsx workbook
that mirrors the tab structure Sean already builds by hand (a filtered
Cadency tab, a filtered SAP/bank tab, and a raw/unfiltered tab per source),
but leads with a "Results" tab: the automated daily compare that his manual
version doesn't have. That's the actual improvement being pitched -- the
familiar tabs are there so the output still looks/feels like what he
already trusts.
"""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl.styles import PatternFill

from .models import ReconStatus

_STATUS_FILLS = {
    ReconStatus.RECONCILED.value: PatternFill("solid", fgColor="D4EDDA"),
    ReconStatus.DOES_NOT_RECONCILE.value: PatternFill("solid", fgColor="F8D7DA"),
    ReconStatus.MANUAL_REVIEW.value: PatternFill("solid", fgColor="FFF3CD"),
}


def export_workbook(sheets: list[tuple[str, pd.DataFrame]]) -> bytes:
    """Write an ordered list of (sheet name, dataframe) pairs to an
    in-memory .xlsx workbook and return the raw bytes (ready for
    st.download_button). Sheet names are truncated to Excel's 31-char limit.

    The first sheet is treated as the results/compare sheet: if it has a
    "Status" column, each row is fully color-filled to match the app's
    RECONCILED / DOES NOT RECONCILE / MANUAL REVIEW convention, so the
    workbook is legible on its own, without the app.
    """
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, df in sheets:
            df.to_excel(writer, sheet_name=name[:31], index=False)

        results_name, results_df = sheets[0]
        if "Status" in results_df.columns:
            ws = writer.sheets[results_name[:31]]
            n_cols = len(results_df.columns)
            for row_idx, status in enumerate(results_df["Status"], start=2):
                fill = _STATUS_FILLS.get(status)
                if fill is not None:
                    for col_idx in range(1, n_cols + 1):
                        ws.cell(row=row_idx, column=col_idx).fill = fill

    return buffer.getvalue()
